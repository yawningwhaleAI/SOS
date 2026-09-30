#!/usr/bin/env python3
"""Step 5 (preview) — price summary + ladders from the derived tables.

Writes an Excel workbook with:
  - SKU_Master     : canonical SKUs and their parsed attributes
  - Price_Summary  : per SKU median/min/max SP, median discount, ₹/100 pulls,
                     ₹/100 ply-sheets, obs count, platforms seen
  - Price_Ladder   : per category, SKUs ranked by median ₹/100 ply-sheets

Always medians + spread, never means (spec §5). Requires openpyxl.

Usage: python analysis/pricing.py
"""
from __future__ import annotations

import json
import sqlite3
import statistics
from collections import defaultdict
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / "data" / "sos.db"
EXPORTS = ROOT / "exports"


def med(v):
    v = [x for x in v if x is not None]
    return round(statistics.median(v), 2) if v else None


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    skus = {r["canonical_sku_id"]: dict(r)
            for r in conn.execute("SELECT * FROM canonical_skus")}
    price = [dict(r) for r in conn.execute("SELECT * FROM price_observations")]
    conn.close()
    if not price:
        print("No price_observations. Run build_canonical.py first.")
        return 0

    # aggregate per SKU
    agg = defaultdict(lambda: defaultdict(list))
    platforms = defaultdict(set)
    for p in price:
        k = p["canonical_sku_id"]
        for f in ("selling_price", "discount_pct", "price_per_100_pulls",
                  "price_per_100_ply_sheets"):
            if p[f] is not None:
                agg[k][f].append(p[f])
        platforms[k].add(p["platform"])

    EXPORTS.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # --- SKU_Master ---
    ws = wb.active
    ws.title = "SKU_Master"
    cols = ["canonical_sku_id", "brand", "category", "ply", "units_per_pack",
            "pulls_per_unit", "total_pulls", "material", "claims", "attr_confidence"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for sid, s in sorted(skus.items(), key=lambda kv: (kv[1]["category"] or "", kv[1]["brand"] or "")):
        claims = ", ".join(json.loads(s["claims"] or "[]"))
        ws.append([sid, s["brand"], s["category"], s["ply"], s["units_per_pack"],
                   s["pulls_per_unit"], s["total_pulls"], s["material"], claims, s["attr_confidence"]])

    # --- Price_Summary ---
    ws2 = wb.create_sheet("Price_Summary")
    hdr = ["canonical_sku_id", "brand", "category", "n_obs", "platforms",
           "median_SP", "min_SP", "max_SP", "median_discount_pct",
           "median_per_100_pulls", "median_per_100_ply_sheets"]
    ws2.append(hdr)
    for c in ws2[1]:
        c.font = Font(bold=True)
    summary = {}
    for k, a in agg.items():
        s = skus.get(k, {})
        sp = a["selling_price"]
        row = [k, s.get("brand"), s.get("category"), len(sp),
               ", ".join(sorted(platforms[k])),
               med(sp), min(sp) if sp else None, max(sp) if sp else None,
               med(a["discount_pct"]), med(a["price_per_100_pulls"]),
               med(a["price_per_100_ply_sheets"])]
        summary[k] = row
        ws2.append(row)

    # --- Price_Ladder (per category, by median ₹/100 ply-sheets, fallback pulls) ---
    ws3 = wb.create_sheet("Price_Ladder")
    ws3.append(["category", "rank", "brand", "canonical_sku_id", "total_pulls",
                "median_per_100_ply_sheets", "median_per_100_pulls", "n_obs"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    by_cat = defaultdict(list)
    for k, a in agg.items():
        s = skus.get(k, {})
        key_metric = med(a["price_per_100_ply_sheets"]) or med(a["price_per_100_pulls"])
        if key_metric is None:
            continue
        by_cat[s.get("category")].append(
            (key_metric, s.get("brand"), k, s.get("total_pulls"),
             med(a["price_per_100_ply_sheets"]), med(a["price_per_100_pulls"]),
             len(a["selling_price"])))
    for cat in sorted(by_cat):
        for rank, row in enumerate(sorted(by_cat[cat]), 1):
            ws3.append([cat, rank, row[1], row[2], row[3], row[4], row[5], row[6]])

    out = EXPORTS / f"SOS_Price_Ladders_{date.today().isoformat()}.xlsx"
    wb.save(out)
    print(f"Wrote {out}")
    print(f"  SKU_Master rows: {len(skus)}")
    print(f"  Price_Summary rows: {len(summary)}")
    print(f"  Price_Ladder categories: {len(by_cat)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
