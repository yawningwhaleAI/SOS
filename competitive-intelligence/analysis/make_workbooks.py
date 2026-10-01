#!/usr/bin/env python3
"""Generate the pricing/positioning deliverables from the cleaned DB.

Outputs to exports/:
  SOS_Brand_Product_Table_<date>.xlsx  (flat table + benchmarks + simulator)
  SOS_Conclusion_Matrix_<date>.xlsx    (one-row-per-category decision table)

Uses clean rows only (is_clean=1). Run after build_canonical.py + clean_data.py.
Single-tenant reference; port/generalise for the product.
"""
from __future__ import annotations
import sqlite3, statistics as st, re, json
from collections import defaultdict, Counter
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / "data" / "sos.db"
EXPORTS = ROOT / "exports"
CORE = ["Facial", "Kitchen towel", "Toilet", "Napkin", "Pocket"]

ALIAS = [(r"honest home", "Honest Home"), (r"10.?on", "10on"), (r"premier", "Premier"),
         (r"luvlap", "Luvlap"), (r"origami", "Origami"), (r"beco", "Beco"),
         (r"wintex", "Wintex"), (r"paseo", "Paseo"), (r"ginni", "GINNI"), (r"kressa", "Kressa")]
def bclean(b):
    n = (b or "").strip().lower()
    for p, name in ALIAS:
        if re.search(p, n): return name
    return (b or "Unknown").strip().title()
def med(v): return round(st.median(v), 1) if v else None
def avg(v): return round(st.mean(v), 1) if v else None
def pctl(v, q):
    if not v: return None
    v = sorted(v); i = min(len(v) - 1, max(0, round(q * (len(v) - 1)))); return round(v[i], 1)

HEAD = Font(bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="1F3A5F")
TITLE = Font(bold=True, size=14); SUB = Font(bold=True, size=11, color="1F3A5F")
SOSF = PatternFill("solid", fgColor="FFF2CC")
def rupee(c): c.number_format = '₹#,##0'
def pct(c): c.number_format = '0%'
def hdr(ws):
    for cc in ws[1]: cc.font = HEAD; cc.fill = HFILL


def load_skus(conn):
    conn.row_factory = sqlite3.Row
    skus = {r['canonical_sku_id']: dict(r) for r in conn.execute("SELECT * FROM canonical_skus")}
    rows = conn.execute("""SELECT p.canonical_sku_id id,p.selling_price sp,p.mrp,p.discount_pct disc,
        p.price_per_100_pulls pp,p.price_per_100_ply_sheets pps,r.product_name_raw name,r.platform,r.url,
        m.rating rt,m.review_count rc
        FROM price_observations p JOIN raw_observations r ON r.obs_id=p.obs_id
        LEFT JOIN market_observations m ON m.obs_id=p.obs_id
        WHERE p.is_clean=1 AND p.selling_price IS NOT NULL""").fetchall()
    agg = defaultdict(lambda: {'sp': [], 'mrp': [], 'disc': [], 'pp': [], 'pps': [], 'rt': [], 'rc': [], 'name': [], 'plat': set(), 'url': None})
    for r in rows:
        a = agg[r['id']]; a['sp'].append(r['sp'])
        for k, f in [('mrp', 'mrp'), ('disc', 'disc'), ('pp', 'pp'), ('pps', 'pps'), ('rt', 'rt')]:
            if r[f] is not None: a[k].append(r[f])
        if r['rc'] is not None: a['rc'].append(r['rc'])
        if r['name']: a['name'].append(r['name'])
        a['plat'].add(r['platform'])
        if r['url'] and not a['url']: a['url'] = r['url']
    SKU = {}
    for sid, a in agg.items():
        s = skus.get(sid, {})
        SKU[sid] = dict(brand=bclean(s.get('brand')), cat=s.get('category') or "", ply=s.get('ply'),
            units=s.get('units_per_pack'), total=s.get('total_pulls'), material=s.get('material') or "not_stated",
            claims=", ".join(json.loads(s.get('claims') or "[]")),
            name=Counter(a['name']).most_common(1)[0][0] if a['name'] else sid,
            mrp=med(a['mrp']), sp=med(a['sp']), disc=med(a['disc']), pp=med(a['pp']), pps=med(a['pps']),
            rt=med(a['rt']), rc=max(a['rc']) if a['rc'] else 0, plat=",".join(sorted(a['plat'])), url=a['url'] or "")
    return SKU


def build_table_workbook(SKU):
    wb = Workbook()
    # Sheet 1: flat table
    ws = wb.active; ws.title = "Brand_Product_Table"
    cols = ["Brand", "Product", "Category", "Ply", "Pack (units)", "Total sheets", "MRP", "Selling price",
            "Discount %", "₹/100 sheets", "Rating", "Reviews", "Material", "Claims", "Platforms", "Product URL"]
    ws.append(cols)
    for s in sorted(SKU.values(), key=lambda x: (x['cat'], x['brand'], -(x['rc'] or 0))):
        ws.append([s['brand'], s['name'][:70], s['cat'], s['ply'], s['units'], s['total'], s['mrp'], s['sp'],
                   round(s['disc'], 2) if s['disc'] is not None else None, s['pp'], s['rt'], s['rc'],
                   s['material'], s['claims'], s['plat'], s['url']])
    hdr(ws); ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:P{ws.max_row}"
    for r in ws.iter_rows(min_row=2):
        rupee(r[6]); rupee(r[7]); pct(r[8])
        u = r[15]
        if isinstance(u.value, str) and u.value.startswith("http"):
            u.hyperlink = u.value; u.font = Font(color="0563C1", underline="single")
    for i, w in enumerate([15, 42, 13, 5, 11, 11, 8, 11, 10, 12, 7, 8, 12, 24, 16, 46], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet 2: benchmarks
    ws = wb.create_sheet("Pricing_Benchmarks")
    ws.append(["Breakdown", "Category", "Group", "#SKUs", "Avg SP", "Median SP", "P25 SP", "P75 SP",
               "Median ₹/100 sheets", "Avg ₹/100 sheets", "Median discount", "Median rating", "Median reviews"])
    def block(dim, keyfn, labelfn):
        g = defaultdict(list)
        for s in SKU.values():
            if not s['cat']: continue
            k = keyfn(s)
            if k is None: continue
            g[k].append(s)
        out = []
        for k, items in g.items():
            sps = [i['sp'] for i in items if i['sp']]
            if not sps: continue
            cat, grp = labelfn(k)
            out.append([dim, cat, grp, len(sps), avg(sps), med(sps), pctl(sps, .25), pctl(sps, .75),
                med([i['pp'] for i in items if i['pp']]), avg([i['pp'] for i in items if i['pp']]),
                round(med([i['disc'] for i in items if i['disc'] is not None]), 2) if any(i['disc'] is not None for i in items) else None,
                med([i['rt'] for i in items if i['rt']]), int(med([i['rc'] for i in items]))])
        return sorted(out, key=lambda r: (str(r[1]), str(r[2])))
    for row in block("By category", lambda s: s['cat'], lambda k: (k, "all")): ws.append(row)
    for row in block("By category × ply", lambda s: (s['cat'], s['ply']), lambda k: (k[0], f"{k[1]}-ply" if k[1] else "ply not stated")): ws.append(row)
    for row in block("By category × material", lambda s: (s['cat'], s['material']), lambda k: (k[0], k[1])): ws.append(row)
    for row in block("By category × pack", lambda s: (s['cat'], s['units']), lambda k: (k[0], f"{k[1]}-pack" if k[1] else "?")): ws.append(row)
    for row in block("By material (all)", lambda s: s['material'], lambda k: ("(all categories)", k)): ws.append(row)
    hdr(ws); ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:M{ws.max_row}"
    for r in ws.iter_rows(min_row=2):
        for ci in (4, 5, 6, 7): rupee(r[ci])
        pct(r[10])
    for i, w in enumerate([22, 14, 18, 7, 9, 10, 9, 9, 18, 16, 13, 12, 13], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # bands for simulator
    band = defaultdict(list)
    for s in SKU.values(): band[(s['cat'], s['units'])].append(s)
    bandref = {}
    for (cat, u), items in band.items():
        sps = [i['sp'] for i in items if i['sp']]
        if len(sps) >= 2: bandref[(cat, u)] = (pctl(sps, .25), med(sps), pctl(sps, .75))

    # Sheet 3: simulator
    ws = wb.create_sheet("SOS_Pricing")
    ws["A1"] = "S.O.S. PRICING SIMULATOR"; ws["A1"].font = TITLE
    ws["A3"] = "Assumptions (edit):"; ws["A3"].font = SUB
    ws["A4"] = "GST rate"; ws["B4"] = 0.18; ws["B4"].number_format = '0%'
    ws["A5"] = "Platform take (commission+fees)"; ws["B5"] = 0.30; ws["B5"].number_format = '0%'
    ws["A7"] = "Brand realisation = SP ÷ (1+GST) × (1−platform take). Fill the yellow cost cells."; ws["A7"].font = Font(italic=True)
    h = ["S.O.S. SKU", "Category", "MRP", "Selling price (SP)", "Cost per pack (FILL)", "Brand realisation",
         "Gross margin ₹", "Gross margin %", "Discount shown", "Band P25", "Band median", "Band P75", "Position vs band"]
    for j, x in enumerate(h, 1):
        cc = ws.cell(9, j, x); cc.font = HEAD; cc.fill = HFILL
    heroes = [("Kitchen towel 2×60", "Kitchen towel", 249, 189, ("Kitchen towel", 2)),
              ("Facial 100", "Facial", 199, 159, ("Facial", 1)),
              ("Facial 200 (2×100)", "Facial", 279, 219, ("Facial", 2)),
              ("Pocket 10×10", "Pocket", 199, 169, ("Pocket", 10)),
              ("Toilet 6-roll 3-ply", "Toilet", 399, 289, ("Toilet", 6))]
    for i, (name, cat, mrp, sp, key) in enumerate(heroes):
        r = 10 + i; p25, m, p75 = bandref.get(key, (None, None, None))
        ws.cell(r, 1, name); ws.cell(r, 2, cat); ws.cell(r, 3, mrp); ws.cell(r, 4, sp)
        ws.cell(r, 6, f"=D{r}/(1+$B$4)*(1-$B$5)"); ws.cell(r, 7, f"=F{r}-E{r}")
        ws.cell(r, 8, f"=IF(F{r}>0,(F{r}-E{r})/F{r},\"\")"); ws.cell(r, 9, f"=IF(C{r}>0,1-D{r}/C{r},\"\")")
        ws.cell(r, 10, p25); ws.cell(r, 11, m); ws.cell(r, 12, p75)
        ws.cell(r, 13, f'=IF(D{r}<J{r},"below band",IF(D{r}>L{r},"above band","within band"))')
        for ci in (3, 4, 5, 6, 7, 10, 11, 12): rupee(ws.cell(r, ci))
        pct(ws.cell(r, 8)); pct(ws.cell(r, 9)); ws.cell(r, 5).fill = SOSF
    ws.column_dimensions['A'].width = 22
    for col in "BCDEFGHIJKLM": ws.column_dimensions[col].width = 15
    return wb, bandref


def build_conclusion(SKU, bandref):
    cat = defaultdict(lambda: {'sp': [], 'pp': [], 'ply': Counter(), 'brev': defaultdict(int), 'prem': []})
    for s in SKU.values():
        if not s['cat']: continue
        if s['sp']: cat[s['cat']]['sp'].append(s['sp'])
        if s['pp']: cat[s['cat']]['pp'].append(s['pp'])
        if s['ply']: cat[s['cat']]['ply'][s['ply']] += 1
        cat[s['cat']]['brev'][s['brand']] += (s['rc'] or 0)
        cat[s['cat']]['prem'].append((s['sp'] or 0, s['brand'], s['rc'] or 0))
    wb = Workbook(); ws = wb.active; ws.title = "Conclusion_Matrix"
    cols = ["Category", "Typical price (median SP)", "Price band (P25–P75)", "₹/100 sheets", "Usual ply",
            "Market leader (reviews)", "Premium that sells", "S.O.S. suggested entry", "Positioning angle"]
    ws.append(cols)
    entry = {"Facial": ("₹149–169", "Design-premium 2-ply"),
             "Kitchen towel": ("₹179–199", "Match Honest Home's sticky premium"),
             "Toilet": ("₹269–299", "3-ply + design = top-of-band hero"),
             "Napkin": ("₹99–129", "Design / occasion premium"),
             "Pocket": ("₹149–179", "Pocket premium sells (Beco proves it)")}
    for ca in CORE:
        d = cat[ca]; sps = d['sp']
        if not sps: continue
        leader = max(d['brev'], key=d['brev'].get)
        ply = f"{d['ply'].most_common(1)[0][0]}-ply" if d['ply'] else "n/a"
        prem = [x for x in sorted(d['prem'], reverse=True) if x[2] >= 500]
        ps = f"{prem[0][1]} ₹{int(prem[0][0])}" if prem else "none sustained"
        ent, ang = entry.get(ca, ("", ""))
        ws.append([ca, f"₹{int(med(sps))}", f"₹{pctl(sps,.25):.0f}–{pctl(sps,.75):.0f}",
                   f"₹{int(med(d['pp']))}" if d['pp'] else "-", ply, f"{leader} ({d['brev'][leader]:,})", ps, ent, ang])
    hdr(ws)
    for cc in ws[1]: cc.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    thin = Side(style="thin", color="D0D0D0"); B = Border(thin, thin, thin, thin)
    for r in ws.iter_rows():
        for c in r: c.border = B; c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 8).fill = SOSF; ws.cell(r, 8).font = Font(bold=True, color="7F6000")
    ws.freeze_panes = "A2"
    for i, w in enumerate([15, 14, 14, 11, 9, 22, 18, 16, 40], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(1, ws.max_row + 1): ws.row_dimensions[r].height = 40
    nr = ws.max_row + 2
    ws.cell(nr, 1, "Read: Origami owns volume everywhere. Premium pricing works — but only with a review moat. Enter near each band's P75, earn reviews, then climb.").font = Font(italic=True, color="555555")
    ws.cell(nr + 1, 1, "Source: cleaned Delhi-NCR quick-commerce snapshot. Prices = medians. Directional, not field-verified. Multipack variants may be under-captured by the scraper.").font = Font(italic=True, size=9, color="888888")
    return wb


def main():
    EXPORTS.mkdir(exist_ok=True)
    conn = sqlite3.connect(DB)
    SKU = load_skus(conn); conn.close()
    today = date.today().isoformat()
    wb1, bandref = build_table_workbook(SKU)
    p1 = EXPORTS / f"SOS_Brand_Product_Table_{today}.xlsx"; wb1.save(p1)
    wb2 = build_conclusion(SKU, bandref)
    p2 = EXPORTS / f"SOS_Conclusion_Matrix_{today}.xlsx"; wb2.save(p2)
    print("wrote", p1); print("wrote", p2); print("SKUs:", len(SKU))


if __name__ == "__main__":
    main()
